from __future__ import annotations

import csv
import hashlib
import io
from datetime import datetime
from typing import Any

import polars as pl
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.connectors.base import RawRecord
from app.config import get_settings
from app.models import SourceRun, SourceSpec, utcnow
from app.services.ingestion import ingest_record
from app.services.normalization import parse_date


ALIASES = {
    "brand_name": ["brand_name", "trade_name", "dba", "business_name", "facility_name", "establishment", "name"],
    "legal_name": ["legal_name", "entity_name", "licensee", "permittee", "company_name"],
    "entity_external_id": ["entity_external_id", "sos_id", "sosid", "entity_id"],
    "entity_type": ["entity_type", "organization_type", "company_type"],
    "entity_status": ["entity_status", "business_status"],
    "registered_agent_name": ["registered_agent_name", "registered_agent", "agent_name"],
    "category": ["category", "business_type", "facility_type", "primary_category"],
    "vertical": ["vertical", "industry_vertical"],
    "naics": ["naics", "naics_code"],
    "website": ["website", "url", "web_site"],
    "phone": ["phone", "telephone", "business_phone"],
    "email": ["email", "business_email"],
    "address1": ["address1", "address", "street", "street_address", "physical_address"],
    "city": ["city", "locality"],
    "state": ["state", "region"],
    "postal_code": ["postal_code", "zip", "zipcode"],
    "county": ["county"],
    "latitude": ["latitude", "lat"],
    "longitude": ["longitude", "lon", "lng"],
    "operating_status": ["operating_status", "status"],
    "person_name": ["person_name", "owner_name", "official_name", "qualifier"],
    "person_role": ["person_role", "role", "title"],
    "ownership_status": ["ownership_status"],
    "permit_kind": ["permit_kind", "permit_type", "license_type", "classification"],
    "permit_number": ["permit_number", "license_number", "file_number"],
    "permit_status": ["permit_status", "license_status"],
    "permit_issued_on": ["permit_issued_on", "issued_on", "issue_date"],
    "permit_expires_on": ["permit_expires_on", "expires_on", "expiration", "expiration_date"],
    "coverage_type": ["coverage_type"],
    "carrier": ["carrier", "insurer"],
    "policy_identifier": ["policy_identifier", "policy_number"],
    "coverage_effective_on": ["coverage_effective_on", "policy_effective_date"],
    "coverage_expires_on": ["coverage_expires_on", "policy_expiration_date"],
    "coverage_as_of": ["coverage_as_of"],
    "trigger_kind": ["trigger_kind"],
    "trigger_date": ["trigger_date", "inspection_date", "event_date"],
    "trigger_label": ["trigger_label"],
}

DATE_FIELDS = {"permit_issued_on", "permit_expires_on", "coverage_effective_on", "coverage_expires_on", "coverage_as_of", "trigger_date", "formed_on"}


def _canonical_headers(headers: list[str], explicit: dict[str, str] | None = None) -> dict[str, str]:
    lowered = {header.strip().lower(): header for header in headers if header}
    mapping: dict[str, str] = {}
    for target, candidates in ALIASES.items():
        if explicit and target in explicit and explicit[target] in headers:
            mapping[target] = explicit[target]
            continue
        for candidate in candidates:
            if candidate in lowered:
                mapping[target] = lowered[candidate]
                break
    return mapping


def _rows_from_file(content: bytes, filename: str) -> list[dict[str, Any]]:
    if filename.lower().endswith(".xls"):
        return pl.read_excel(io.BytesIO(content), engine="calamine").to_dicts()
    if filename.lower().endswith((".xlsx", ".xlsm")):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        return [dict(zip(headers, values)) for values in rows]
    try:
        return pl.read_csv(
            io.BytesIO(content), infer_schema_length=10_000, ignore_errors=True,
            encoding="utf8-lossy",
        ).to_dicts()
    except pl.exceptions.PolarsError:
        text = content.decode("utf-8-sig", errors="replace")
        return list(csv.DictReader(io.StringIO(text)))


def import_tabular(
    db: Session,
    source: SourceSpec,
    content: bytes,
    filename: str,
    explicit_mapping: dict[str, str] | None = None,
) -> SourceRun:
    digest = hashlib.sha256(content).hexdigest()
    suffix = ".xls" if filename.lower().endswith(".xls") else (".xlsx" if filename.lower().endswith((".xlsx", ".xlsm")) else ".csv")
    raw_dir = get_settings().data_dir / "raw" / source.id
    raw_dir.mkdir(parents=True, exist_ok=True)
    snapshot_path = raw_dir / f"{digest}{suffix}"
    if not snapshot_path.exists():
        snapshot_path.write_bytes(content)
    rows = _rows_from_file(content, filename)
    run = SourceRun(
        source_id=source.id, status="running", started_at=utcnow(),
        scope={"filename": filename, "content_hash": digest, "snapshot_path": str(snapshot_path)},
    )
    db.add(run)
    db.flush()
    if not rows:
        run.status = "completed"
        run.finished_at = utcnow()
        db.commit()
        return run
    mapping = _canonical_headers(list(rows[0].keys()), explicit_mapping)
    for index, row in enumerate(rows, start=2):
        values = {target: row.get(header) for target, header in mapping.items()}
        for field in DATE_FIELDS:
            if field in values:
                values[field] = parse_date(values[field])
        for numeric in ("latitude", "longitude"):
            if values.get(numeric) not in (None, ""):
                try:
                    values[numeric] = float(values[numeric])
                except (TypeError, ValueError):
                    values[numeric] = None
        for field, value in list(values.items()):
            if field not in DATE_FIELDS | {"latitude", "longitude"} and value not in (None, "") and not isinstance(value, str):
                values[field] = str(value)
        values["source_record_id"] = str(row.get("source_record_id") or row.get("id") or f"{filename}:{index}")
        values["source_url"] = source.base_url
        values["state"] = (str(values.get("state") or source.jurisdiction).upper()[:2] if source.jurisdiction in {"NC", "SC", "VA"} or values.get("state") else None)
        if source.id in {"osha", "epa_echo", "sc_food_grades", "nc_food_inspections"}:
            values["trigger_risk_review_only"] = True
            values["trigger_kind"] = values.get("trigger_kind") or "regulatory_review"
            values["trigger_label"] = values.get("trigger_label") or f"{source.name} record"
        record = RawRecord(**{key: value for key, value in values.items() if value not in (None, "")})
        run.records_seen += 1
        account, created = ingest_record(db, source, record, run)
        if account:
            run.records_created += int(created)
            run.records_updated += int(not created)
    run.status = "completed"
    run.finished_at = utcnow()
    source.last_success_at = utcnow()
    db.commit()
    return run
